"""Safe file transformation for the local ``/masking`` operation.

Only formats that can be rewritten without flattening their structure are
accepted: UTF text/Markdown, CSV, and XLSX.  Sources are resolved inside an
explicitly authorized personal/project (or caller-supplied temporary) root,
never modified, and outputs are staged in the source directory before a final
no-clobber publication step.  A source hash is checked again immediately
before publication to close the common read/transform/write race.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import inspect
import io
import os
import re
import stat
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path, PurePath
from typing import Any, Awaitable, Callable, Iterable, Mapping, Sequence
from uuid import UUID

from .privacy_masking_boundary import PrivacyMaskingError

SUPPORTED_MASKING_EXTENSIONS = frozenset({".txt", ".md", ".csv", ".xlsx"})
TEXT_MASKING_EXTENSIONS = frozenset({".txt", ".md", ".csv"})
DEFAULT_MAX_FILE_BYTES = 50 * 1024 * 1024
_SAFE_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,127}$")


class MaskingFileError(PrivacyMaskingError):
    """Base failure for unsafe, unsupported, or unpublishable files."""


class MaskingAttachmentAuthorizationError(MaskingFileError):
    """The requested source is outside the caller's authorized storage scope."""


class MaskingSourceChangedError(MaskingFileError):
    """The source changed between staging and publication."""


@dataclass(frozen=True)
class AuthorizedMaskingAttachment:
    """Resolved source descriptor; path fields are never serialized publicly."""

    path: Path = field(repr=False)
    display_name: str = field(default="", repr=False)
    scope: str = ""

    @property
    def extension(self) -> str:
        return self.path.suffix.casefold()


@dataclass(frozen=True)
class MaskedFileResult:
    """Published masked file metadata.

    ``source_path`` and ``output_path`` are useful to the local delivery layer
    but are excluded from repr/public projections so assistant content cannot
    accidentally include a raw workspace path.  ``source_name`` is likewise
    internal; callers should show only ``output_name`` to the user.
    """

    output_name: str
    source_sha256: str
    output_sha256: str
    size_bytes: int
    output_path: Path = field(repr=False, compare=False)
    source_path: Path = field(repr=False, compare=False)
    source_name: str = field(default="", repr=False, compare=False)
    extension: str = ""

    @property
    def path(self) -> Path:
        return self.output_path

    def to_public_dict(self) -> dict[str, Any]:
        """Return attachment metadata safe for an assistant/result payload."""

        # Do not include source path/name.  A generated name is intentionally
        # retained for download UX (``example_masked.xlsx``) without exposing
        # where the user's original file resides.
        return {
            "name": self.output_name,
            "filename": self.output_name,
            "kind": "masked_file",
            "size_bytes": int(self.size_bytes),
            "sha256": self.output_sha256,
            "extension": self.extension,
        }

    as_dict = to_public_dict


@dataclass
class _StagedFile:
    attachment: AuthorizedMaskingAttachment
    source_sha256: str
    output_path: Path
    temporary_path: Path
    data: bytes
    findings: list[Any]
    semantic_status: str
    source_name: str


def _is_link_or_reparse(path: Path) -> bool:
    try:
        info = os.lstat(path)
    except FileNotFoundError:
        return False
    except OSError:
        return True
    return stat.S_ISLNK(info.st_mode) or bool(
        getattr(info, "st_file_attributes", 0)
        & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
    )


def _assert_no_links(path: Path, *, stop: Path | None = None) -> None:
    """Reject symlink/reparse components from ``stop`` through ``path``."""

    candidate = Path(os.path.abspath(os.fspath(path)))
    stop_abs = (
        Path(os.path.abspath(os.fspath(stop)))
        if stop is not None
        else Path(candidate.anchor or os.sep)
    )
    try:
        relative_parts = candidate.relative_to(stop_abs).parts
    except ValueError:
        relative_parts = candidate.parts
        stop_abs = Path(candidate.anchor or os.sep)
    current = stop_abs
    if _is_link_or_reparse(current):
        raise MaskingAttachmentAuthorizationError("storage root is a link/reparse point")
    for part in relative_parts:
        current = current / part
        if _is_link_or_reparse(current):
            raise MaskingAttachmentAuthorizationError(
                "attachment path contains a link/reparse point"
            )


def _safe_identifier(value: Any, label: str) -> str:
    raw = str(value or "").strip()
    if not raw or not _SAFE_ID_RE.fullmatch(raw):
        raise MaskingAttachmentAuthorizationError(f"invalid {label}")
    return raw


def _workspace_root(value: str | os.PathLike[str] | None) -> Path:
    if value is not None:
        root = Path(value).expanduser()
    else:
        # Import lazily to keep the service usable in stripped/unit-test builds.
        try:
            from ..tools.file_explorer.storage_context import get_base_storage_dir

            root = get_base_storage_dir()
        except Exception:
            root = Path(os.environ.get("AOITALK_WORKSPACES_DIR", "./workspaces"))
    root = Path(os.path.abspath(os.fspath(root)))
    if _is_link_or_reparse(root):
        raise MaskingAttachmentAuthorizationError("workspace root is a link/reparse point")
    return root.resolve(strict=False)


def _scope_roots(
    *,
    root: Path,
    user_id: Any = None,
    project_id: Any = None,
) -> list[tuple[str, Path]]:
    roots: list[tuple[str, Path]] = []
    if user_id:
        uid = _safe_identifier(user_id, "user id")
        roots.append(("personal", root / "_users" / f"user_{uid}"))
    if project_id:
        pid = _safe_identifier(project_id, "project id")
        roots.append(("project", root / "_projects" / f"project_{pid}"))
    return roots


def _contained(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True


def _raw_attachment_value(item: Any) -> tuple[str, str]:
    if isinstance(item, (str, os.PathLike)):
        raw = os.fspath(item)
        return str(raw), Path(str(raw)).name
    # The authenticated chat resolver commonly returns ``(metadata,
    # normalized_path)`` pairs.  Accept that internal shape while still
    # requiring the final path to pass the same scope/link checks below.
    if isinstance(item, (tuple, list)) and len(item) == 2 and isinstance(item[0], Mapping):
        metadata = dict(item[0])
        metadata.setdefault("path", item[1])
        item = metadata
    if not isinstance(item, Mapping):
        raise MaskingAttachmentAuthorizationError("attachment descriptor is malformed")
    raw = ""
    for key in (
        "project_relative_path",
        "relative_path",
        "path",
        "source_path",
        "file_path",
        "temp_path",
        "absolute_path",
    ):
        candidate = item.get(key)
        if candidate not in (None, ""):
            raw = str(candidate).strip()
            break
    display = str(
        item.get("name")
        or item.get("filename")
        or item.get("display_name")
        or (Path(raw).name if raw else "")
    ).strip()
    if not raw:
        raise MaskingAttachmentAuthorizationError("attachment path is required")
    return raw, display


def resolve_authorized_masking_attachment(
    item: Any,
    *,
    user_id: Any = None,
    project_id: Any = None,
    workspace_root: str | os.PathLike[str] | None = None,
    authorized_roots: Iterable[str | os.PathLike[str]] = (),
    authorization_callback: Callable[..., Any] | None = None,
) -> AuthorizedMaskingAttachment:
    """Resolve one attachment while enforcing personal/project containment.

    ``authorization_callback`` is an optional application ACL seam.  It may
    return a path (or ``True`` to continue with the descriptor path); the final
    containment/link checks still run regardless of the callback result.
    """

    raw, display_name = _raw_attachment_value(item)
    root = _workspace_root(workspace_root)
    scope_roots = _scope_roots(root=root, user_id=user_id, project_id=project_id)
    extra_roots: list[Path] = []
    for value in authorized_roots:
        # Check the caller-declared root before resolving it.  Resolving first
        # would erase a symlink/reparse component and make an otherwise
        # untrusted temporary root look like a regular in-scope directory.
        declared_root = Path(os.path.abspath(os.fspath(value)))
        root_anchor = Path(declared_root.anchor or os.sep)
        _assert_no_links(declared_root, stop=root_anchor)
        resolved_root = declared_root.resolve(strict=False)
        _assert_no_links(
            resolved_root,
            stop=Path(resolved_root.anchor or os.sep),
        )
        extra_roots.append(resolved_root)
    allowed = scope_roots + [("authorized", value) for value in extra_roots]
    if not allowed:
        raise MaskingAttachmentAuthorizationError(
            "an authenticated personal/project scope is required"
        )

    callback_value: Any = None
    if authorization_callback is not None:
        if not callable(authorization_callback):
            raise MaskingAttachmentAuthorizationError("attachment authorization callback is invalid")
        try:
            try:
                callback_value = authorization_callback(
                    item,
                    user_id=user_id,
                    project_id=project_id,
                )
            except TypeError:
                callback_value = authorization_callback(item)
            if inspect.isawaitable(callback_value):
                raise MaskingAttachmentAuthorizationError(
                    "async authorization callback requires resolve_async"
                )
        except MaskingAttachmentAuthorizationError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise MaskingAttachmentAuthorizationError("attachment authorization failed") from exc
        if callback_value is False or callback_value is None:
            raise MaskingAttachmentAuthorizationError("attachment is not authorized")
        if isinstance(callback_value, (str, os.PathLike)):
            raw = os.fspath(callback_value)

    normalized = str(raw).strip().replace("\\", "/")
    if not normalized:
        raise MaskingAttachmentAuthorizationError("attachment path is required")
    if "\x00" in normalized:
        raise MaskingAttachmentAuthorizationError("attachment path is malformed")

    # Map canonical storage-relative paths into the configured workspace root,
    # checking the owner component before any filesystem resolution.
    canonical = normalized.lstrip("/")
    canonical_parts = PurePath(canonical).parts
    candidate: Path
    declared_scope = ""
    if canonical_parts and canonical_parts[0] == "_users":
        if len(canonical_parts) < 3 or not canonical_parts[1].startswith("user_"):
            raise MaskingAttachmentAuthorizationError("personal attachment path is malformed")
        owner = canonical_parts[1][len("user_") :]
        if not user_id or owner.casefold() != str(user_id).strip().casefold():
            raise MaskingAttachmentAuthorizationError("personal attachment is outside the authorized user")
        declared_scope = "personal"
        candidate = root.joinpath(*canonical_parts)
    elif canonical_parts and canonical_parts[0] == "_projects":
        if len(canonical_parts) < 3 or not canonical_parts[1].startswith("project_"):
            raise MaskingAttachmentAuthorizationError("project attachment path is malformed")
        owner = canonical_parts[1][len("project_") :]
        if not project_id or owner.casefold() != str(project_id).strip().casefold():
            raise MaskingAttachmentAuthorizationError("project attachment is outside the authorized project")
        declared_scope = "project"
        candidate = root.joinpath(*canonical_parts)
    elif Path(normalized).is_absolute() or re.match(r"^[A-Za-z]:/", normalized):
        candidate = Path(normalized)
    else:
        # Relative paths are interpreted in the project first, then personal
        # scope.  A project is explicit and therefore takes precedence.
        preferred = next((path for scope, path in scope_roots if scope == "project"), None)
        if preferred is None:
            preferred = next((path for scope, path in scope_roots if scope == "personal"), None)
        if preferred is None:
            # A caller-supplied authorized root is valid for temporary uploads.
            preferred = extra_roots[0] if extra_roots else None
            declared_scope = "authorized"
        if preferred is None:
            raise MaskingAttachmentAuthorizationError("attachment scope is missing")
        candidate = preferred.joinpath(*PurePath(normalized).parts)

    candidate_abs = Path(os.path.abspath(os.fspath(candidate)))
    # Lexical traversal is rejected before ``resolve`` so a missing path cannot
    # smuggle ``..`` through a later symlink/reparse check.
    if any(part == ".." for part in PurePath(normalized).parts):
        raise MaskingAttachmentAuthorizationError("attachment traversal is not allowed")
    resolved = candidate_abs.resolve(strict=False)
    matched_scope = declared_scope
    matched_root: Path | None = None
    for scope, allowed_root in allowed:
        allowed_resolved = allowed_root.resolve(strict=False)
        if _contained(resolved, allowed_resolved):
            matched_scope = matched_scope or scope
            matched_root = allowed_resolved
            break
    if matched_root is None:
        raise MaskingAttachmentAuthorizationError("attachment is outside the authorized storage scope")
    # Check the lexical path before resolving it.  Resolving first erases a
    # symlink/reparse component from the path, so a link that points to another
    # file *inside* the authorized root would otherwise look indistinguishable
    # from a regular file and be accepted.  Re-check the resolved path as well
    # to fail closed on any link introduced by a concurrently changed parent.
    _assert_no_links(candidate_abs, stop=matched_root)
    _assert_no_links(resolved, stop=matched_root)
    if not resolved.is_file():
        raise MaskingAttachmentAuthorizationError("attachment file was not found")
    if _is_link_or_reparse(resolved):
        raise MaskingAttachmentAuthorizationError("attachment file is a link/reparse point")
    suffix = resolved.suffix.casefold()
    if suffix not in SUPPORTED_MASKING_EXTENSIONS:
        raise MaskingFileError(
            f"unsupported masking file type: {suffix or '(none)'}"
        )
    # A display name is UX-only.  Never allow it to redirect the source or
    # manufacture a path; output naming uses the resolved basename.
    safe_display = PurePath(display_name.replace("\\", "/")).name.strip() if display_name else resolved.name
    if not safe_display or safe_display in {".", ".."}:
        safe_display = resolved.name
    return AuthorizedMaskingAttachment(
        path=resolved,
        display_name=safe_display,
        scope=matched_scope,
    )


# Common shorter name used by route/integration code.
resolve_masking_attachment = resolve_authorized_masking_attachment


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _decode_text(data: bytes, path: Path) -> tuple[str, str]:
    try:
        from ..tools.text_content import decode_text_bytes

        return decode_text_bytes(
            data,
            path=path,
            known_text_extensions=TEXT_MASKING_EXTENSIONS,
        )
    except ImportError:
        for encoding in ("utf-8-sig", "utf-16", "cp932"):
            try:
                return data.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        raise MaskingFileError("text file encoding is not supported")
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("text file is not safely readable") from exc


def _masker_accepts_source_kind(masker: Callable[..., Any]) -> bool:
    try:
        signature = inspect.signature(masker)
        if any(parameter.kind == parameter.VAR_KEYWORD for parameter in signature.parameters.values()):
            return True
        return "source_kind" in signature.parameters
    except (TypeError, ValueError):
        # Unknown callable signatures receive the richer call first.  A
        # TypeError from such a callable is treated as its real failure rather
        # than retried, preventing duplicate sidecar requests.
        return True


def _call_masker(masker: Callable[..., Any], text: str, source_kind: str) -> Any:
    result = (
        masker(text, source_kind=source_kind)
        if _masker_accepts_source_kind(masker)
        else masker(text)
    )
    if inspect.isawaitable(result):
        raise MaskingFileError("async masker requires transform_async")
    return result


async def _call_masker_async(masker: Callable[..., Any], text: str, source_kind: str) -> Any:
    result = (
        masker(text, source_kind=source_kind)
        if _masker_accepts_source_kind(masker)
        else masker(text)
    )
    if inspect.isawaitable(result):
        result = await result
    return result


def _masked_value(result: Any) -> tuple[str, list[Any], str]:
    if isinstance(result, str):
        return result, [], "disabled"
    def _field(name: str, default: Any = None) -> Any:
        if isinstance(result, Mapping):
            return result.get(name, default)
        return getattr(result, name, default)

    value = _field("value")
    if value is None:
        value = _field("masked_text")
    if value is None:
        value = _field("final_payload")
    if value is None:
        value = _field("payload")
    if not isinstance(value, str):
        raise MaskingFileError("masking callback returned no text projection")
    findings = list(_field("findings", ()) or ())
    semantic_status = str(
        _field("semantic_status", "disabled") or "disabled"
    ).strip().casefold()
    if semantic_status in {"failed", "error", "unavailable"}:
        raise MaskingFileError("semantic privacy masking failed")
    return value, findings, semantic_status


def _formula_surface(value: str) -> str:
    """Return an XLSX formula skeleton with string-literal bodies elided.

    Formula references/operators must survive masking unchanged.  Literal
    text inside double quotes is the one safe surface to replace (for
    example ``="a@example.com"``).  A semantic detector that labels ``A1``
    in ``=A1&"name"`` must therefore fail closed instead of publishing an
    invalid or silently retargeted workbook.  Excel escapes a quote inside a
    string as ``""``; those pairs are consumed as literal content here rather
    than being mistaken for a new formula surface boundary.
    """

    if not isinstance(value, str):
        return ""
    surface: list[str] = []
    index = 0
    length = len(value)
    while index < length:
        char = value[index]
        if char != '"':
            surface.append(char)
            index += 1
            continue

        surface.append('"')
        index += 1
        # Collapse the whole quoted body to one sentinel so a replacement
        # may change its length without changing the surrounding formula.
        surface.append(chr(0))
        while index < length:
            if value[index] != '"':
                index += 1
                continue
            if index + 1 < length and value[index + 1] == '"':
                # Escaped quote inside the same Excel string literal.
                index += 2
                continue
            surface.append('"')
            index += 1
            break
        else:
            # Preserve malformed/unclosed quote topology; the workbook
            # re-open verification will reject any serialization failure.
            surface.append(chr(1))
    return "".join(surface)


def _xlsx_header_footer_has_text(item: Any) -> bool:
    for section_name in ("left", "center", "right"):
        section = getattr(item, section_name, None)
        if str(getattr(section, "text", "") or ""):
            return True
    return False


def _reject_unsupported_xlsx_surfaces(workbook: Any) -> None:
    """Fail closed for OOXML text surfaces not covered by the masker.

    ``openpyxl`` rewrites these structures without giving the masking boundary
    a safe, lossless way to traverse every text-bearing field.  Publishing a
    workbook with an untouched comment/header/validation/etc. could therefore
    leave a secret in the derived file.  Reject the whole workbook instead of
    silently copying an unmasked surface.
    """

    try:
        if len(getattr(workbook, "defined_names", ()) or ()):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    except TypeError:
        # A malformed/future workbook object is not safe to project.
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")

    properties = getattr(workbook, "properties", None)
    for name in (
        "title",
        "subject",
        "description",
        "identifier",
        "language",
        "lastModifiedBy",
        "category",
        "contentStatus",
        "version",
        "revision",
        "keywords",
    ):
        if str(getattr(properties, name, "") or "").strip():
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    creator = str(getattr(properties, "creator", "") or "").strip()
    if creator and creator.casefold() != "openpyxl":
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    if getattr(workbook, "custom_doc_props", None):
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    if getattr(workbook, "_number_formats", None):
        # Custom number formats are serialized as literal OOXML text and are
        # outside the cell-value masking walk.  Refuse them rather than copy a
        # potentially sensitive format string into the derived workbook.
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    if getattr(workbook, "code_name", None) or getattr(workbook, "chartsheets", None):
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
    if getattr(workbook, "_pivots", None):
        raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")

    header_footer_names = (
        "oddHeader",
        "evenHeader",
        "firstHeader",
        "oddFooter",
        "evenFooter",
        "firstFooter",
    )
    for sheet in getattr(workbook, "worksheets", ()):
        if str(getattr(getattr(sheet, "sheet_properties", None), "codeName", "") or "").strip():
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        cells = tuple(getattr(sheet, "_cells", {}).values())
        for cell in cells:
            if getattr(cell, "comment", None) is not None:
                raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
            if getattr(cell, "hyperlink", None) is not None:
                raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        for name in header_footer_names:
            if _xlsx_header_footer_has_text(getattr(sheet, name, None)):
                raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        data_validations = getattr(getattr(sheet, "data_validations", None), "dataValidation", ())
        if data_validations:
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        try:
            if list(getattr(sheet, "conditional_formatting", ())):
                raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        except TypeError:
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        if getattr(sheet, "tables", None):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        if getattr(sheet, "_charts", None) or getattr(sheet, "_images", None):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        if getattr(sheet, "_pivots", None):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        if getattr(sheet, "legacy_drawing", None):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        if getattr(sheet, "_rels", None):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        auto_filter = getattr(sheet, "auto_filter", None)
        if getattr(auto_filter, "filterColumn", None) or getattr(
            auto_filter, "sortState", None
        ):
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")
        scenarios = getattr(getattr(sheet, "scenarios", None), "scenario", ())
        if scenarios:
            raise MaskingFileError("XLSX contains unsupported text-bearing surfaces")


def _validate_xlsx_style_text_sync(
    workbook: Any,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> None:
    """Check preserved style labels/font names through the masking boundary."""

    values: list[str] = []
    for style in getattr(workbook, "_named_styles", ()):
        name = str(getattr(style, "name", "") or "")
        if name:
            values.append(name)
    for font in getattr(workbook, "_fonts", ()):
        name = str(getattr(font, "name", "") or "")
        if name:
            values.append(name)
    seen: set[str] = set()
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        masked, value_findings, status = _masked_value(
            _call_masker(masker, value, "masking_xlsx_style_text")
        )
        if masked != value:
            raise MaskingFileError("XLSX style text requires masking")
        findings.extend(value_findings)
        semantic_states.append(status)


def _iter_xlsx_xml_strings(data: bytes) -> Iterable[str]:
    """Yield every textual XML node/attribute in an XLSX package.

    The openpyxl object model does not expose every OOXML part (for example
    threaded comments, pivot caches, and custom extensions).  A final package
    scan gives the trusted boundary a format-independent check before any
    bytes are published.  Non-XML package parts are rejected because they are
    not traversable by the masking boundary.
    """

    try:
        archive = zipfile.ZipFile(io.BytesIO(data), "r")
    except (OSError, zipfile.BadZipFile) as exc:
        raise MaskingFileError("XLSX package could not be inspected safely") from exc
    try:
        names = [info.filename for info in archive.infolist()]
        if len(names) != len(set(names)):
            raise MaskingFileError("XLSX package contains duplicate parts")
        for info in archive.infolist():
            name = str(info.filename or "")
            lower_name = name.casefold()
            if not lower_name.endswith((".xml", ".rels")):
                raise MaskingFileError("XLSX contains unsupported package surfaces")
            try:
                root = ET.fromstring(archive.read(info))
            except (OSError, ET.ParseError) as exc:
                raise MaskingFileError("XLSX package contains invalid XML") from exc
            for element in root.iter():
                if not isinstance(element.tag, str):
                    if element.text:
                        yield str(element.text)
                    continue
                if element.text:
                    yield str(element.text)
                if element.tail:
                    yield str(element.tail)
                for attribute in element.attrib.values():
                    if attribute:
                        yield str(attribute)
    finally:
        archive.close()


def _verify_xlsx_ooxml_text_sync(
    data: bytes,
    *,
    masker: Callable[..., Any],
) -> None:
    """Reject any serialized OOXML string that still requires masking."""

    seen: set[str] = set()
    for value in _iter_xlsx_xml_strings(data):
        if not value or not value.strip() or value in seen:
            continue
        seen.add(value)
        masked, _, _ = _masked_value(
            _call_masker(masker, value, "masking_xlsx_ooxml")
        )
        if masked != value:
            raise MaskingFileError("XLSX output contains unmasked text")


async def _verify_xlsx_ooxml_text_async(
    data: bytes,
    *,
    masker: Callable[..., Any],
) -> None:
    """Async counterpart of :func:`_verify_xlsx_ooxml_text_sync`."""

    seen: set[str] = set()
    for value in _iter_xlsx_xml_strings(data):
        if not value or not value.strip() or value in seen:
            continue
        seen.add(value)
        masked, _, _ = _masked_value(
            await _call_masker_async(masker, value, "masking_xlsx_ooxml")
        )
        if masked != value:
            raise MaskingFileError("XLSX output contains unmasked text")


async def _validate_xlsx_style_text_async(
    workbook: Any,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> None:
    """Async counterpart of :func:`_validate_xlsx_style_text_sync`."""

    values: list[str] = []
    for style in getattr(workbook, "_named_styles", ()):
        name = str(getattr(style, "name", "") or "")
        if name:
            values.append(name)
    for font in getattr(workbook, "_fonts", ()):
        name = str(getattr(font, "name", "") or "")
        if name:
            values.append(name)
    seen: set[str] = set()
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        masked, value_findings, status = _masked_value(
            await _call_masker_async(masker, value, "masking_xlsx_style_text")
        )
        if masked != value:
            raise MaskingFileError("XLSX style text requires masking")
        findings.extend(value_findings)
        semantic_states.append(status)


def _validate_xlsx_sheet_titles_sync(
    workbook: Any,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> None:
    """Run sheet names through the same boundary without mutating them."""

    for sheet in getattr(workbook, "worksheets", ()):
        title = str(getattr(sheet, "title", "") or "")
        if not title:
            raise MaskingFileError("XLSX sheet title is invalid")
        masked, title_findings, status = _masked_value(
            _call_masker(masker, title, "masking_xlsx_sheet_title")
        )
        if masked != title:
            raise MaskingFileError("XLSX sheet title requires masking")
        findings.extend(title_findings)
        semantic_states.append(status)


async def _validate_xlsx_sheet_titles_async(
    workbook: Any,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> None:
    """Async counterpart of :func:`_validate_xlsx_sheet_titles_sync`."""

    for sheet in getattr(workbook, "worksheets", ()):
        title = str(getattr(sheet, "title", "") or "")
        if not title:
            raise MaskingFileError("XLSX sheet title is invalid")
        masked, title_findings, status = _masked_value(
            await _call_masker_async(masker, title, "masking_xlsx_sheet_title")
        )
        if masked != title:
            raise MaskingFileError("XLSX sheet title requires masking")
        findings.extend(title_findings)
        semantic_states.append(status)


def _csv_dialect(text: str) -> csv.Dialect:
    sample = text[:64 * 1024]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def _transform_csv_sync(
    text: str,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> str:
    dialect = _csv_dialect(text)
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True))
    except csv.Error as exc:
        raise MaskingFileError("CSV structure is invalid") from exc
    out_rows: list[list[str]] = []
    for row in rows:
        out: list[str] = []
        for cell in row:
            if cell == "":
                out.append(cell)
                continue
            value, cell_findings, status = _masked_value(
                _call_masker(masker, cell, "masking_csv_cell")
            )
            out.append(value)
            findings.extend(cell_findings)
            semantic_states.append(status)
        out_rows.append(out)
    output = io.StringIO(newline="")
    writer = csv.writer(
        output,
        dialect=dialect,
        lineterminator=getattr(dialect, "lineterminator", "\r\n"),
    )
    writer.writerows(out_rows)
    return output.getvalue()


async def _transform_csv_async(
    text: str,
    *,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> str:
    dialect = _csv_dialect(text)
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True))
    except csv.Error as exc:
        raise MaskingFileError("CSV structure is invalid") from exc
    out_rows: list[list[str]] = []
    for row in rows:
        out: list[str] = []
        for cell in row:
            if cell == "":
                out.append(cell)
                continue
            value, cell_findings, status = _masked_value(
                await _call_masker_async(masker, cell, "masking_csv_cell")
            )
            out.append(value)
            findings.extend(cell_findings)
            semantic_states.append(status)
        out_rows.append(out)
    output = io.StringIO(newline="")
    writer = csv.writer(
        output,
        dialect=dialect,
        lineterminator=getattr(dialect, "lineterminator", "\r\n"),
    )
    writer.writerows(out_rows)
    return output.getvalue()


def _transform_xlsx_sync(
    data: bytes,
    *,
    source_name: str,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> bytes:
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell
    except Exception as exc:  # pragma: no cover - dependency is required by pyproject
        raise MaskingFileError("XLSX masking support is unavailable") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("XLSX workbook could not be opened safely") from exc
    # External links can trigger provider/network resolution in downstream
    # tooling and cannot be proven safe to preserve here.
    if getattr(workbook, "_external_links", None):
        raise MaskingFileError("XLSX external links are not supported")
    if getattr(workbook, "vba_archive", None) is not None:
        raise MaskingFileError("macro-enabled XLSX surfaces are not supported")
    try:
        _reject_unsupported_xlsx_surfaces(workbook)
        _validate_xlsx_sheet_titles_sync(
            workbook,
            masker=masker,
            findings=findings,
            semantic_states=semantic_states,
        )
        _validate_xlsx_style_text_sync(
            workbook,
            masker=masker,
            findings=findings,
            semantic_states=semantic_states,
        )
        for sheet in workbook.worksheets:
            # Iterating existing cells (rather than ``iter_rows`` over the full
            # dimensions) avoids materializing enormous sparse worksheets.
            for cell in tuple(sheet._cells.values()):  # type: ignore[attr-defined]
                if isinstance(cell, MergedCell):
                    continue
                value = cell.value
                if not isinstance(value, str) or not value:
                    continue
                source_kind = (
                    "masking_xlsx_formula"
                    if value.startswith("=")
                    else "masking_xlsx_cell"
                )
                masked, cell_findings, status = _masked_value(
                    _call_masker(masker, value, source_kind)
                )
                if value.startswith("=") and (
                    not masked.startswith("=")
                    or _formula_surface(masked) != _formula_surface(value)
                ):
                    # A formula is a structured surface.  If a detector tries
                    # to replace the whole expression rather than only a
                    # literal inside it, refusing publication is safer than
                    # silently flattening/corrupting the workbook.
                    raise MaskingFileError(
                        "XLSX formula masking would change the formula surface"
                    )
                cell.value = masked
                findings.extend(cell_findings)
                semantic_states.append(status)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    except MaskingFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("XLSX workbook could not be safely rewritten") from exc
    finally:
        try:
            workbook.close()
        except Exception:
            pass


async def _transform_xlsx_async(
    data: bytes,
    *,
    source_name: str,
    masker: Callable[..., Any],
    findings: list[Any],
    semantic_states: list[str],
) -> bytes:
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell
    except Exception as exc:  # pragma: no cover
        raise MaskingFileError("XLSX masking support is unavailable") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("XLSX workbook could not be opened safely") from exc
    if getattr(workbook, "_external_links", None):
        raise MaskingFileError("XLSX external links are not supported")
    if getattr(workbook, "vba_archive", None) is not None:
        raise MaskingFileError("macro-enabled XLSX surfaces are not supported")
    try:
        _reject_unsupported_xlsx_surfaces(workbook)
        await _validate_xlsx_sheet_titles_async(
            workbook,
            masker=masker,
            findings=findings,
            semantic_states=semantic_states,
        )
        await _validate_xlsx_style_text_async(
            workbook,
            masker=masker,
            findings=findings,
            semantic_states=semantic_states,
        )
        for sheet in workbook.worksheets:
            for cell in tuple(sheet._cells.values()):  # type: ignore[attr-defined]
                if isinstance(cell, MergedCell):
                    continue
                value = cell.value
                if not isinstance(value, str) or not value:
                    continue
                source_kind = (
                    "masking_xlsx_formula"
                    if value.startswith("=")
                    else "masking_xlsx_cell"
                )
                masked, cell_findings, status = _masked_value(
                    await _call_masker_async(masker, value, source_kind)
                )
                if value.startswith("=") and (
                    not masked.startswith("=")
                    or _formula_surface(masked) != _formula_surface(value)
                ):
                    raise MaskingFileError(
                        "XLSX formula masking would change the formula surface"
                    )
                cell.value = masked
                findings.extend(cell_findings)
                semantic_states.append(status)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    except MaskingFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("XLSX workbook could not be safely rewritten") from exc
    finally:
        try:
            workbook.close()
        except Exception:
            pass


def _output_candidate(source: Path, index: int = 0) -> Path:
    suffix = source.suffix
    stem = source.name[: -len(suffix)] if suffix else source.name
    if index <= 0:
        name = f"{stem}_masked{suffix}"
    else:
        name = f"{stem}_masked_{index + 1}{suffix}"
    return source.with_name(name)


def _write_staging_file(data: bytes, parent: Path, name: str) -> Path:
    fd, raw_path = tempfile.mkstemp(prefix=f".{name}.", suffix=".masking.tmp", dir=str(parent))
    temp_path = Path(raw_path)
    try:
        with os.fdopen(fd, "wb") as stream:
            fd = -1
            stream.write(data)
            stream.flush()
            try:
                os.fsync(stream.fileno())
            except OSError:
                pass
    except Exception:
        if fd >= 0:
            os.close(fd)
        temp_path.unlink(missing_ok=True)
        raise
    return temp_path


def _choose_available_output(source: Path) -> Path:
    for index in range(0, 10_000):
        candidate = _output_candidate(source, index)
        if not candidate.exists() and not _is_link_or_reparse(candidate):
            return candidate
    raise MaskingFileError("could not reserve a no-clobber masking output name")


def _publish_no_clobber(temp_path: Path, candidate: Path) -> Path:
    """Publish without replacing an existing output, tolerating a race."""

    current = candidate
    # ``candidate`` may already be a numbered name selected because an earlier
    # ``*_masked`` output exists.  Keep the original source stem so a race on
    # that numbered candidate advances monotonically (``foo_masked_3``), rather
    # than producing ``foo_masked_2_masked_2``.
    suffix = candidate.suffix
    candidate_stem = candidate.name[: -len(suffix)] if suffix else candidate.name
    source_stem = candidate_stem
    match = re.fullmatch(r"(?P<base>.+)_masked(?:_(?P<index>\d+))?", candidate_stem)
    if match:
        source_stem = match.group("base")
        next_index = int(match.group("index") or "1") + 1
    else:
        next_index = 2
    for index in range(0, 10_000):
        try:
            # A hard link gives us an atomic create-if-absent operation on the
            # same filesystem; the temp is then unlinked.  It never clobbers a
            # path another process won between staging and publication.
            os.link(temp_path, current)
            temp_path.unlink(missing_ok=True)
            return current
        except FileExistsError:
            current = candidate.with_name(
                f"{source_stem}_masked_{next_index + index}{suffix}"
            )
            continue
        except OSError as exc:
            # There is no portable rename-no-replace primitive.  Falling back
            # to ``exists()`` + ``os.replace()`` would introduce a TOCTOU
            # window in which a concurrent writer could be clobbered.  The
            # trusted masking boundary therefore fails closed when the local
            # filesystem cannot provide the hard-link create-if-absent
            # primitive instead of weakening the no-clobber guarantee.
            raise MaskingFileError(
                "filesystem cannot provide atomic no-clobber masking output"
            ) from exc
    raise MaskingFileError("could not publish a no-clobber masking output")


def _verify_xlsx(data: bytes, expected_sheetnames: Sequence[str]) -> None:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        try:
            if list(workbook.sheetnames) != list(expected_sheetnames):
                raise MaskingFileError("masked XLSX sheet structure changed")
            if getattr(workbook, "_external_links", None):
                raise MaskingFileError("masked XLSX contains unsupported external links")
        finally:
            workbook.close()
    except MaskingFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MaskingFileError("masked XLSX could not be reopened") from exc


class MaskingFileTransformer:
    """Stage and publish one or more authorized masking file outputs."""

    def __init__(
        self,
        *,
        workspace_root: str | os.PathLike[str] | None = None,
        user_id: Any = None,
        project_id: Any = None,
        authorized_roots: Iterable[str | os.PathLike[str]] = (),
        authorization_callback: Callable[..., Any] | None = None,
        max_file_bytes: int = DEFAULT_MAX_FILE_BYTES,
    ) -> None:
        self.workspace_root = workspace_root
        self.user_id = user_id
        self.project_id = project_id
        self.authorized_roots = tuple(authorized_roots)
        self.authorization_callback = authorization_callback
        try:
            self.max_file_bytes = max(1, int(max_file_bytes))
        except (TypeError, ValueError, OverflowError):
            self.max_file_bytes = DEFAULT_MAX_FILE_BYTES

    def resolve_attachment(self, item: Any) -> AuthorizedMaskingAttachment:
        return resolve_authorized_masking_attachment(
            item,
            user_id=self.user_id,
            project_id=self.project_id,
            workspace_root=self.workspace_root,
            authorized_roots=self.authorized_roots,
            authorization_callback=self.authorization_callback,
        )

    def _read_source(self, attachment: AuthorizedMaskingAttachment) -> tuple[bytes, str]:
        source = attachment.path
        try:
            stat_result = source.stat()
        except OSError as exc:
            raise MaskingFileError("masking source could not be inspected") from exc
        if stat_result.st_size > self.max_file_bytes:
            raise MaskingFileError("masking source exceeds the configured size limit")
        try:
            data = source.read_bytes()
        except OSError as exc:
            raise MaskingFileError("masking source could not be read") from exc
        if len(data) != stat_result.st_size:
            raise MaskingSourceChangedError("masking source changed while reading")
        return data, _sha256(data)

    def transform(
        self,
        attachments: Iterable[Any],
        *,
        masker: Callable[..., Any],
    ) -> tuple[tuple[MaskedFileResult, ...], tuple[Any, ...], str]:
        """Synchronously transform all files, staging before publication."""

        items = list(attachments or ())
        if not items:
            return (), (), "disabled"
        # Re-run authorization even for callers that hand us an
        # ``AuthorizedMaskingAttachment`` instance.  The dataclass is public
        # and can be forged by an embedding/client; trusting its type would
        # let an arbitrary absolute path bypass the owner/project containment
        # checks at this boundary.
        resolved = [
            self.resolve_attachment(
                item.path if isinstance(item, AuthorizedMaskingAttachment) else item
            )
            for item in items
        ]
        staged: list[_StagedFile] = []
        published_paths: list[Path] = []
        all_findings: list[Any] = []
        semantic_states: list[str] = []
        try:
            for attachment in resolved:
                data, source_hash = self._read_source(attachment)
                findings: list[Any] = []
                states: list[str] = []
                if attachment.extension in TEXT_MASKING_EXTENSIONS:
                    text, encoding = _decode_text(data, attachment.path)
                    if attachment.extension == ".csv":
                        masked_text = _transform_csv_sync(text, masker=masker, findings=findings, semantic_states=states)
                    else:
                        masked_text, one_findings, status = _masked_value(
                            _call_masker(masker, text, "masking_text_file")
                        )
                        findings.extend(one_findings)
                        states.append(status)
                    try:
                        output_data = masked_text.encode(encoding)
                    except (LookupError, UnicodeEncodeError) as exc:
                        raise MaskingFileError("masked text could not be encoded") from exc
                elif attachment.extension == ".xlsx":
                    # Capture sheet names before transformation for reopen check.
                    try:
                        from openpyxl import load_workbook

                        check_workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
                        sheetnames = tuple(check_workbook.sheetnames)
                        check_workbook.close()
                    except Exception as exc:  # noqa: BLE001
                        raise MaskingFileError("XLSX workbook could not be inspected") from exc
                    output_data = _transform_xlsx_sync(
                        data,
                        source_name=attachment.path.name,
                        masker=masker,
                        findings=findings,
                        semantic_states=states,
                    )
                    _verify_xlsx(output_data, sheetnames)
                    _verify_xlsx_ooxml_text_sync(output_data, masker=masker)
                else:  # resolver currently rejects this; keep a defensive guard
                    raise MaskingFileError("unsupported masking file type")
                output_path = _choose_available_output(attachment.path)
                temporary_path = _write_staging_file(output_data, attachment.path.parent, output_path.name)
                staged.append(
                    _StagedFile(
                        attachment=attachment,
                        source_sha256=source_hash,
                        output_path=output_path,
                        temporary_path=temporary_path,
                        data=output_data,
                        findings=findings,
                        semantic_status=(states[-1] if states else "disabled"),
                        source_name=attachment.path.name,
                    )
                )
                all_findings.extend(findings)
                semantic_states.extend(states)

            # Reopen/re-hash every source before publishing any staged output.
            for entry in staged:
                current = entry.attachment.path.read_bytes()
                if _sha256(current) != entry.source_sha256:
                    raise MaskingSourceChangedError("masking source changed before publication")

            published: list[MaskedFileResult] = []
            for entry in staged:
                published_path = _publish_no_clobber(entry.temporary_path, entry.output_path)
                published_paths.append(published_path)
                output_bytes = published_path.read_bytes()
                if output_bytes != entry.data:
                    raise MaskingFileError("published masking output verification failed")
                output_hash = _sha256(output_bytes)
                published.append(
                    MaskedFileResult(
                        output_name=published_path.name,
                        source_sha256=entry.source_sha256,
                        output_sha256=output_hash,
                        size_bytes=len(output_bytes),
                        output_path=published_path,
                        source_path=entry.attachment.path,
                        source_name=entry.source_name,
                        extension=entry.attachment.extension,
                    )
                )
            status = (
                "success"
                if any(
                    str(state).strip().casefold() == "success"
                    for state in semantic_states
                )
                else "disabled"
            )
            return tuple(published), tuple(all_findings), status
        except Exception:
            for entry in staged:
                entry.temporary_path.unlink(missing_ok=True)
            # Outputs are only published after all source checks.  If a failure
            # occurs while publishing, remove only files we just published; a
            # pre-existing `_masked` file is never touched.
            for path in published_paths:
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise

    async def transform_async(
        self,
        attachments: Iterable[Any],
        *,
        masker: Callable[..., Any],
    ) -> tuple[tuple[MaskedFileResult, ...], tuple[Any, ...], str]:
        """Async counterpart sharing one masker/gateway across all files."""

        items = list(attachments or ())
        if not items:
            return (), (), "disabled"
        # Keep the async path subject to the same authorization revalidation
        # as the synchronous transformer; a caller cannot bypass scope checks
        # merely by constructing the public descriptor dataclass.
        resolved = [
            self.resolve_attachment(
                item.path if isinstance(item, AuthorizedMaskingAttachment) else item
            )
            for item in items
        ]
        staged: list[_StagedFile] = []
        published_paths: list[Path] = []
        all_findings: list[Any] = []
        semantic_states: list[str] = []
        try:
            for attachment in resolved:
                data, source_hash = self._read_source(attachment)
                findings: list[Any] = []
                states: list[str] = []
                if attachment.extension in TEXT_MASKING_EXTENSIONS:
                    text, encoding = _decode_text(data, attachment.path)
                    if attachment.extension == ".csv":
                        masked_text = await _transform_csv_async(text, masker=masker, findings=findings, semantic_states=states)
                    else:
                        masked_text, one_findings, status = _masked_value(
                            await _call_masker_async(masker, text, "masking_text_file")
                        )
                        findings.extend(one_findings)
                        states.append(status)
                    try:
                        output_data = masked_text.encode(encoding)
                    except (LookupError, UnicodeEncodeError) as exc:
                        raise MaskingFileError("masked text could not be encoded") from exc
                elif attachment.extension == ".xlsx":
                    try:
                        from openpyxl import load_workbook

                        check_workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
                        sheetnames = tuple(check_workbook.sheetnames)
                        check_workbook.close()
                    except Exception as exc:  # noqa: BLE001
                        raise MaskingFileError("XLSX workbook could not be inspected") from exc
                    output_data = await _transform_xlsx_async(
                        data,
                        source_name=attachment.path.name,
                        masker=masker,
                        findings=findings,
                        semantic_states=states,
                    )
                    _verify_xlsx(output_data, sheetnames)
                    await _verify_xlsx_ooxml_text_async(output_data, masker=masker)
                else:
                    raise MaskingFileError("unsupported masking file type")
                output_path = _choose_available_output(attachment.path)
                temporary_path = _write_staging_file(output_data, attachment.path.parent, output_path.name)
                staged.append(
                    _StagedFile(
                        attachment=attachment,
                        source_sha256=source_hash,
                        output_path=output_path,
                        temporary_path=temporary_path,
                        data=output_data,
                        findings=findings,
                        semantic_status=(states[-1] if states else "disabled"),
                        source_name=attachment.path.name,
                    )
                )
                all_findings.extend(findings)
                semantic_states.extend(states)
            for entry in staged:
                current = entry.attachment.path.read_bytes()
                if _sha256(current) != entry.source_sha256:
                    raise MaskingSourceChangedError("masking source changed before publication")
            published: list[MaskedFileResult] = []
            for entry in staged:
                published_path = _publish_no_clobber(entry.temporary_path, entry.output_path)
                published_paths.append(published_path)
                output_bytes = published_path.read_bytes()
                if output_bytes != entry.data:
                    raise MaskingFileError("published masking output verification failed")
                published.append(
                    MaskedFileResult(
                        output_name=published_path.name,
                        source_sha256=entry.source_sha256,
                        output_sha256=_sha256(output_bytes),
                        size_bytes=len(output_bytes),
                        output_path=published_path,
                        source_path=entry.attachment.path,
                        source_name=entry.source_name,
                        extension=entry.attachment.extension,
                    )
                )
            status = (
                "success"
                if any(
                    str(state).strip().casefold() == "success"
                    for state in semantic_states
                )
                else "disabled"
            )
            return tuple(published), tuple(all_findings), status
        except Exception:
            for entry in staged:
                entry.temporary_path.unlink(missing_ok=True)
            for path in published_paths:
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
            raise

    transform_files = transform
    transform_files_async = transform_async


__all__ = [
    "AuthorizedMaskingAttachment",
    "DEFAULT_MAX_FILE_BYTES",
    "MaskingAttachmentAuthorizationError",
    "MaskingFileError",
    "MaskingFileTransformer",
    "MaskingSourceChangedError",
    "MaskedFileResult",
    "SUPPORTED_MASKING_EXTENSIONS",
    "resolve_authorized_masking_attachment",
    "resolve_masking_attachment",
]
